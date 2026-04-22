"""Workbook scoring: cell values, formulas, structure, formula-drift."""
from __future__ import annotations

from openpyxl import Workbook as OpenpyxlWorkbook

from evals.scoring import MAX_DIFFS_PER_CASE, score
from shotcut.spreadsheet.workbook import Workbook


def _wb(build) -> Workbook:
    pyxl = OpenpyxlWorkbook()
    build(pyxl)
    return Workbook(pyxl)


def test_identical_workbooks_pass() -> None:
    def build(wb):
        ws = wb.active
        ws.title = "S"
        ws["A1"] = 42
        ws["A2"] = "=A1*2"

    report = score(_wb(build), _wb(build))
    assert report.passed
    assert report.diffs == []


def test_cell_value_mismatch_records_diff() -> None:
    def expected(wb):
        wb.active.title = "S"
        wb.active["A1"] = 100

    def actual(wb):
        wb.active.title = "S"
        wb.active["A1"] = 999

    report = score(_wb(expected), _wb(actual))
    assert not report.passed
    assert len(report.diffs) == 1
    diff = report.diffs[0]
    assert diff.axis == "cell_value"
    assert diff.location == "S!A1"
    assert diff.expected == 100
    assert diff.actual == 999


def test_formula_mismatch_but_same_evaluation_is_formula_drift() -> None:
    """=B1+B2 vs =SUM(B1:B2) evaluate identically → labeled drift, not fail."""
    def expected(wb):
        wb.active.title = "S"
        wb.active["B1"] = 10
        wb.active["B2"] = 20
        wb.active["C1"] = "=SUM(B1:B2)"

    def actual(wb):
        wb.active.title = "S"
        wb.active["B1"] = 10
        wb.active["B2"] = 20
        wb.active["C1"] = "=B1+B2"

    report = score(_wb(expected), _wb(actual))
    # A drift diff is present but its note indicates same evaluated value.
    assert len(report.diffs) == 1
    diff = report.diffs[0]
    assert diff.axis == "formula"
    assert diff.note and "drift" in diff.note


def test_formula_mismatch_with_different_evaluation_is_failure() -> None:
    def expected(wb):
        wb.active.title = "S"
        wb.active["A1"] = 10
        wb.active["A2"] = "=A1*2"  # → 20

    def actual(wb):
        wb.active.title = "S"
        wb.active["A1"] = 10
        wb.active["A2"] = "=A1*3"  # → 30

    report = score(_wb(expected), _wb(actual))
    assert not report.passed
    diff = report.diffs[0]
    assert diff.axis == "formula"
    assert diff.note is None  # no drift note when evals differ


def test_formula_vs_literal_recorded() -> None:
    def expected(wb):
        wb.active.title = "S"
        wb.active["A1"] = "=SUM(1,2)"

    def actual(wb):
        wb.active.title = "S"
        wb.active["A1"] = 3

    report = score(_wb(expected), _wb(actual))
    assert len(report.diffs) == 1
    assert report.diffs[0].note and "expected a formula" in report.diffs[0].note


def test_missing_sheet_is_structural_diff() -> None:
    def expected(wb):
        wb.active.title = "Inputs"
        wb.active["A1"] = 1
        wb.create_sheet("Model")

    def actual(wb):
        wb.active.title = "Inputs"
        wb.active["A1"] = 1

    report = score(_wb(expected), _wb(actual))
    assert not report.passed
    structural = [d for d in report.diffs if d.axis == "structure"]
    assert any("Model" in d.location for d in structural)


def test_extra_sheet_is_structural_diff() -> None:
    def expected(wb):
        wb.active.title = "Inputs"
        wb.active["A1"] = 1

    def actual(wb):
        wb.active.title = "Inputs"
        wb.active["A1"] = 1
        wb.create_sheet("Unexpected")

    report = score(_wb(expected), _wb(actual))
    structural = [d for d in report.diffs if d.axis == "structure"]
    assert any("Unexpected" in d.location for d in structural)


def test_numeric_tolerance_allows_float_drift() -> None:
    def expected(wb):
        wb.active.title = "S"
        wb.active["A1"] = 1.0

    def actual(wb):
        wb.active.title = "S"
        wb.active["A1"] = 1.0 + 1e-10  # within 1e-6 tolerance

    assert score(_wb(expected), _wb(actual)).passed


def test_many_diffs_truncated_to_max() -> None:
    """The scorer caps diffs at MAX_DIFFS_PER_CASE and flags the rest."""
    def expected(wb):
        wb.active.title = "S"
        for i in range(1, MAX_DIFFS_PER_CASE + 10):
            wb.active[f"A{i}"] = i

    def actual(wb):
        wb.active.title = "S"
        for i in range(1, MAX_DIFFS_PER_CASE + 10):
            wb.active[f"A{i}"] = 0  # every cell wrong

    report = score(_wb(expected), _wb(actual))
    assert len(report.diffs) == MAX_DIFFS_PER_CASE
    assert report.diffs_truncated is True
