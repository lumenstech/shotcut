"""`Workbook.with_actions_applied` context manager.

Stage 6 shape-ready for an eventual copy-on-write upgrade; current
implementation is a full clone. Tests verify:
- Clone yields a workbook with the actions applied.
- The original workbook is never mutated.
- The clone is closed on exit (no leaked ZipFile handles).
"""
from __future__ import annotations

import gc
from pathlib import Path

import pytest
from openpyxl import Workbook as OpenpyxlWorkbook

from shotcut.spreadsheet.actions import WriteFormula, WriteValue
from shotcut.spreadsheet.workbook import Workbook


def _build(tmp_path: Path) -> Workbook:
    pyxl = OpenpyxlWorkbook()
    pyxl.active.title = "Sheet1"
    pyxl.active["A1"] = 10
    pyxl.active["A2"] = 20
    path = tmp_path / "wb.xlsx"
    pyxl.save(path)
    return Workbook.from_xlsx(path)


def test_clone_has_actions_applied(tmp_path: Path) -> None:
    wb = _build(tmp_path)
    actions = [
        WriteValue(sheet="Sheet1", target="A3", value=30),
        WriteFormula(sheet="Sheet1", target="A4", formula="=SUM(A1:A3)"),
    ]
    with wb.with_actions_applied(actions) as clone:
        assert clone.raw["Sheet1"]["A3"].value == 30
        assert clone.raw["Sheet1"]["A4"].value == "=SUM(A1:A3)"
        assert clone.evaluate_cell("Sheet1", "A4") == 60


def test_original_is_unchanged(tmp_path: Path) -> None:
    wb = _build(tmp_path)
    before = {
        c.coordinate: c.value
        for row in wb.raw["Sheet1"].iter_rows()
        for c in row
        if c.value is not None
    }

    with wb.with_actions_applied(
        [WriteValue(sheet="Sheet1", target="A1", value=999999)]
    ):
        pass

    after = {
        c.coordinate: c.value
        for row in wb.raw["Sheet1"].iter_rows()
        for c in row
        if c.value is not None
    }
    assert before == after


@pytest.mark.filterwarnings("error::pytest.PytestUnraisableExceptionWarning")
def test_clone_does_not_leak_zip_handles(tmp_path: Path) -> None:
    """Stage 6 adds a new clone path — it must satisfy the Stage 4
    ResourceWarning regression guard too."""
    wb = _build(tmp_path)
    for _ in range(100):
        with wb.with_actions_applied(
            [WriteValue(sheet="Sheet1", target="B1", value=1)]
        ):
            pass
    gc.collect()
