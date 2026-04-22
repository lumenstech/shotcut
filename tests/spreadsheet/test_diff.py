"""Labeled diff tests.

Covers every label (AGENT_ADDED / AGENT_OVERWROTE_AGENT /
AGENT_OVERWROTE_USER / USER_MODIFIED defensive assertion) and the
force_override_cells interplay.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook as OpenpyxlWorkbook

from shotcut.spreadsheet.actions import WriteValue
from shotcut.spreadsheet.diff import DiffLabel, UserModificationError, diff
from shotcut.spreadsheet.occupancy import OccupancyMap
from shotcut.spreadsheet.workbook import Workbook


def _save(wb: Workbook, tmp_path: Path, name: str) -> Path:
    path = tmp_path / name
    wb.save(path)
    return path


def _empty_workbook() -> Workbook:
    pyxl = OpenpyxlWorkbook()
    pyxl.active.title = "Sheet1"
    return Workbook(pyxl)


def test_agent_added_label_when_empty_cell_fills(tmp_path: Path):
    """No original upload → every write is agent_added."""
    before = _empty_workbook()
    after = _empty_workbook()
    after.apply(WriteValue(sheet="Sheet1", target="A1", value=42))

    result = diff(before, after, OccupancyMap.empty())
    assert len(result) == 1
    assert result[0].label == DiffLabel.AGENT_ADDED
    assert result[0].before is None
    assert result[0].after == 42


def test_agent_overwrote_agent_when_agent_cell_changes(tmp_path: Path):
    """Both states have a cell; the cell wasn't user-sourced."""
    before = _empty_workbook()
    before.apply(WriteValue(sheet="Sheet1", target="A1", value=10))
    after = _empty_workbook()
    after.apply(WriteValue(sheet="Sheet1", target="A1", value=20))

    result = diff(before, after, OccupancyMap.empty())
    assert len(result) == 1
    assert result[0].label == DiffLabel.AGENT_OVERWROTE_AGENT


def test_agent_overwrote_user_when_force_override_cell_changes(tmp_path: Path):
    """User cell + cell in force_override_cells → AGENT_OVERWROTE_USER."""
    # Build an original with a user-occupied cell.
    pyxl = OpenpyxlWorkbook()
    pyxl.active.title = "Sheet1"
    pyxl.active["A1"] = 1000
    original_path = tmp_path / "original.xlsx"
    pyxl.save(original_path)
    original = Workbook.from_xlsx(original_path)
    occupancy = OccupancyMap.from_original(original)

    # Start from original, mutate.
    before = Workbook.from_xlsx(original_path)
    after = Workbook.from_xlsx(original_path)
    after.apply(WriteValue(sheet="Sheet1", target="A1", value=9999))

    forced = {("Sheet1", "A1")}
    result = diff(before, after, occupancy, force_override_cells=forced)
    assert len(result) == 1
    assert result[0].label == DiffLabel.AGENT_OVERWROTE_USER


def test_user_modified_raises_when_user_cell_changes_without_override(tmp_path: Path):
    """Defensive: user cell changed but not in force_override_cells → assert.

    This label indicates something has bypassed the orchestrator's state
    machine. It should never occur in normal operation; the diff raises
    rather than silently returning the label, so a test failure is loud.
    """
    pyxl = OpenpyxlWorkbook()
    pyxl.active.title = "Sheet1"
    pyxl.active["A1"] = 1000
    original_path = tmp_path / "original.xlsx"
    pyxl.save(original_path)
    original = Workbook.from_xlsx(original_path)
    occupancy = OccupancyMap.from_original(original)

    before = Workbook.from_xlsx(original_path)
    after = Workbook.from_xlsx(original_path)
    after.apply(WriteValue(sheet="Sheet1", target="A1", value=9999))

    with pytest.raises(UserModificationError):
        diff(before, after, occupancy, force_override_cells=set())


def test_empty_diff_when_no_changes(tmp_path: Path):
    before = _empty_workbook()
    before.apply(WriteValue(sheet="Sheet1", target="A1", value=1))
    after = _empty_workbook()
    after.apply(WriteValue(sheet="Sheet1", target="A1", value=1))

    assert diff(before, after, OccupancyMap.empty()) == []


def test_mixed_labels_in_one_diff(tmp_path: Path):
    """Exercise multiple label types in a single diff call."""
    # Original: Sheet1!A1 is user-occupied.
    pyxl = OpenpyxlWorkbook()
    pyxl.active.title = "Sheet1"
    pyxl.active["A1"] = "user"
    original_path = tmp_path / "original.xlsx"
    pyxl.save(original_path)
    occupancy = OccupancyMap.from_original(Workbook.from_xlsx(original_path))

    # Before: original + agent wrote to B1.
    before = Workbook.from_xlsx(original_path)
    before.apply(WriteValue(sheet="Sheet1", target="B1", value=10))
    # After: A1 force-overridden, B1 changed by agent, C1 newly added.
    after = Workbook.from_xlsx(original_path)
    after.apply(WriteValue(sheet="Sheet1", target="A1", value="overridden"))
    after.apply(WriteValue(sheet="Sheet1", target="B1", value=20))
    after.apply(WriteValue(sheet="Sheet1", target="C1", value="new"))

    forced = {("Sheet1", "A1")}
    result = diff(before, after, occupancy, force_override_cells=forced)
    labels_by_ref = {(d.sheet, d.ref): d.label for d in result}
    assert labels_by_ref[("Sheet1", "A1")] == DiffLabel.AGENT_OVERWROTE_USER
    assert labels_by_ref[("Sheet1", "B1")] == DiffLabel.AGENT_OVERWROTE_AGENT
    assert labels_by_ref[("Sheet1", "C1")] == DiffLabel.AGENT_ADDED
