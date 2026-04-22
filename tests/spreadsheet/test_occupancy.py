"""OccupancyMap unit tests.

Covers the is_occupied rule and the check() contract against each action
type. End-to-end occupancy-blocking through the orchestrator is exercised
in test_nondestructive.py.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.styles import Font, PatternFill

from shotcut.spreadsheet.actions import (
    AddSheet,
    FormatCell,
    SetColumnWidth,
    WriteFormula,
    WriteValue,
)
from shotcut.spreadsheet.occupancy import OccupancyMap
from shotcut.spreadsheet.workbook import Workbook


def _build_workbook(tmp_path: Path, build: callable) -> Workbook:
    """Helper: build an xlsx via the callback, save, load as our Workbook."""
    pyxl = OpenpyxlWorkbook()
    build(pyxl)
    path = tmp_path / "wb.xlsx"
    pyxl.save(path)
    return Workbook.from_xlsx(path)


# ---------------------------------------------------------------------------
# is_occupied rule
# ---------------------------------------------------------------------------


def test_empty_map_treats_every_cell_as_free():
    """A session with no uploaded original starts fully free."""
    occupancy = OccupancyMap.empty()
    assert not occupancy.is_user_occupied("Sheet", "A1")
    assert occupancy.user_cell_count == 0


def test_cell_with_value_is_occupied(tmp_path: Path):
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["A1"] = 42

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)
    assert occupancy.is_user_occupied("Sheet1", "A1")
    assert not occupancy.is_user_occupied("Sheet1", "A2")


def test_cell_with_formula_is_occupied(tmp_path: Path):
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["B1"] = "=SUM(1,2)"

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)
    assert occupancy.is_user_occupied("Sheet1", "B1")


def test_cell_with_non_default_format_but_no_value_is_occupied(tmp_path: Path):
    """A yellow fill on an empty cell is still user intent."""
    def build(wb):
        ws = wb.active
        ws.title = "Sheet1"
        yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws["C1"].fill = yellow

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)
    assert occupancy.is_user_occupied("Sheet1", "C1")


def test_cell_with_bold_font_is_occupied(tmp_path: Path):
    def build(wb):
        ws = wb.active
        ws.title = "Sheet1"
        ws["D1"].font = Font(bold=True)

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)
    assert occupancy.is_user_occupied("Sheet1", "D1")


def test_untouched_empty_cell_does_not_block_overwrite(tmp_path: Path):
    """Default format + no value → NOT occupied, writes are free to proceed.

    Acceptance criterion: "Default-format empty cell does not block overwrite."
    """
    def build(wb):
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = 1  # one occupied cell so the workbook isn't empty

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    assert occupancy.is_user_occupied("Sheet1", "A1")
    # E5 was never touched: no value, default font, no fill, no borders.
    assert not occupancy.is_user_occupied("Sheet1", "E5")

    action = WriteValue(sheet="Sheet1", target="E5", value=99)
    assert occupancy.check(action).blocked is False


def test_merged_range_cells_all_occupied(tmp_path: Path):
    """Every cell inside a merge is treated as occupied — including the
    empty shells openpyxl stores for the spanned cells."""
    def build(wb):
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Header"
        ws.merge_cells("A1:C1")

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)
    for ref in ("A1", "B1", "C1"):
        assert occupancy.is_user_occupied("Sheet1", ref), f"{ref} should be occupied"


# ---------------------------------------------------------------------------
# check() contract against each action type
# ---------------------------------------------------------------------------


def test_check_allows_write_to_empty_cell(tmp_path: Path):
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["A1"] = 1

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    action = WriteValue(sheet="Sheet1", target="B1", value=99)
    result = occupancy.check(action)
    assert result.blocked is False
    assert result.overwrites == []


def test_check_blocks_write_to_user_cell(tmp_path: Path):
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["A1"] = 1000

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    action = WriteValue(sheet="Sheet1", target="A1", value=2000)
    result = occupancy.check(action)
    assert result.blocked is True
    assert ("Sheet1", "A1") in result.overwrites
    assert result.reason is not None
    assert "Sheet1!A1" in result.reason


def test_check_blocks_formula_targeting_user_range(tmp_path: Path):
    def build(wb):
        wb.active.title = "Sheet1"
        for i in range(1, 4):
            wb.active[f"A{i}"] = i

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    action = WriteFormula(sheet="Sheet1", target="A1:A3", formula="=ROW()")
    result = occupancy.check(action)
    assert result.blocked is True
    assert sorted(result.overwrites) == [("Sheet1", "A1"), ("Sheet1", "A2"), ("Sheet1", "A3")]


def test_check_blocks_format_on_user_cell(tmp_path: Path):
    """FormatCell must respect occupancy — it rewrites user styling."""
    def build(wb):
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Important"

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    action = FormatCell(sheet="Sheet1", target="A1", bold=True)
    result = occupancy.check(action)
    assert result.blocked is True


def test_check_allows_add_sheet_even_without_occupancy(tmp_path: Path):
    """AddSheet doesn't target cells; occupancy is not a concern."""
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["A1"] = 1

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    result = occupancy.check(AddSheet(sheet="NewSheet"))
    assert result.blocked is False


def test_check_allows_set_column_width(tmp_path: Path):
    """Column width is not occupancy-blocking — it's a layout hint."""
    def build(wb):
        wb.active.title = "Sheet1"
        wb.active["A1"] = 1

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    result = occupancy.check(SetColumnWidth(sheet="Sheet1", target="A1", width=20.0))
    assert result.blocked is False


def test_check_reason_truncates_long_overwrite_lists(tmp_path: Path):
    """More than 5 overwrites in a single action — reason caps the listing."""
    def build(wb):
        wb.active.title = "Sheet1"
        for i in range(1, 11):
            wb.active[f"A{i}"] = i

    wb = _build_workbook(tmp_path, build)
    occupancy = OccupancyMap.from_original(wb)

    action = WriteValue(sheet="Sheet1", target="A1:A10", value=0)
    result = occupancy.check(action)
    assert result.blocked is True
    assert len(result.overwrites) == 10
    assert result.reason is not None
    assert "and 5 more" in result.reason
