"""Curated golden-set cases.

Three cases that CI gates on. Each builds an expected `Workbook`
programmatically via openpyxl — no .xlsx binaries committed. Adding a
case means: implement a builder here, add the `case_id` to
`baseline.json` with verdict `"passed"`, and commit both.

The case_ids here match the keys in `evals/results/baseline.json`;
drift between the two fails the regression gate. This is intentional:
if you want to add or remove a case, it's a two-file commit.
"""
from __future__ import annotations

from collections.abc import Callable

from openpyxl import Workbook as OpenpyxlWorkbook

from evals.runner import EvalCase
from shotcut.spreadsheet.workbook import Workbook


def _wrap(build_fn: Callable[[OpenpyxlWorkbook], None]) -> Workbook:
    pyxl = OpenpyxlWorkbook()
    build_fn(pyxl)
    return Workbook(pyxl)


def _sum_of_column() -> Workbook:
    def build(wb: OpenpyxlWorkbook) -> None:
        ws = wb.active
        ws.title = "Sheet"
        for i in range(1, 6):
            ws[f"A{i}"] = i * 10
        ws["B1"] = "=SUM(A1:A5)"

    return _wrap(build)


def _formula_reference_chain() -> Workbook:
    def build(wb: OpenpyxlWorkbook) -> None:
        ws = wb.active
        ws.title = "Sheet"
        ws["A1"] = 100
        ws["A2"] = "=A1*2"
        ws["A3"] = "=A2+50"
        ws["A4"] = "=A3/5"

    return _wrap(build)


def _multi_sheet_total() -> Workbook:
    def build(wb: OpenpyxlWorkbook) -> None:
        inputs = wb.active
        inputs.title = "Inputs"
        inputs["A1"] = 1000
        inputs["A2"] = 500
        model = wb.create_sheet("Model")
        model["A1"] = "=Inputs!A1+Inputs!A2"

    return _wrap(build)


CASES: list[EvalCase] = [
    EvalCase(
        case_id="sum_of_column",
        prompt="Put the sum of A1:A5 in B1.",
        expected=_sum_of_column(),
    ),
    EvalCase(
        case_id="formula_reference_chain",
        prompt=(
            "Given A1=100, compute a three-step chain: "
            "A2 is A1 doubled, A3 is A2 + 50, A4 is A3 / 5."
        ),
        expected=_formula_reference_chain(),
    ),
    EvalCase(
        case_id="multi_sheet_total",
        prompt=(
            "In a new sheet 'Model', cell A1, sum Inputs!A1 and "
            "Inputs!A2 via a formula."
        ),
        expected=_multi_sheet_total(),
    ),
]
