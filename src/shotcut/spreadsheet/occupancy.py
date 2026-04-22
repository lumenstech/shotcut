"""Cell occupancy — which cells are user-sourced, which are free for agents.

Design (see docs/decisions/0002-schema-consolidation.md and the Stage 3
amendments): provenance is *derived from the path split*, not tracked
separately.

- Cells present in `original.xlsx` → user-sourced
- Cells in `current.xlsx` not in `original.xlsx` → agent-sourced (inferable
  from the workbook alone, not persisted here)
- A cell is **occupied** iff:
    value is not None               (literal or formula — openpyxl stores
                                     formulas as strings starting with '=')
  OR has_non_default_format          (yellow fill on an "empty" cell is
                                     still user intent)
  OR is_part_of_merge                (touching any merged cell is a merge
                                     violation even if the cell itself is
                                     empty-but-spanned)

The orchestrator consults `OccupancyMap.check(action)` before applying;
if blocked it stages the action as `pending_approval` instead of raising.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from shotcut.spreadsheet.actions import (
    Action,
    FormatCell,
    WriteFormula,
    WriteValue,
)

if TYPE_CHECKING:
    from shotcut.spreadsheet.workbook import Workbook


@dataclass(frozen=True)
class OccupancyCheck:
    """Result of checking a proposed action against the map.

    `blocked` is True iff the action would overwrite a user-sourced cell.
    `reason` is a human-readable explanation suitable for
    `Action.approval_required_reason`. `overwrites` is the full list of
    (sheet, ref) conflicts — the reason only cites the first few.
    """

    blocked: bool
    reason: str | None = None
    overwrites: list[tuple[str, str]] = field(default_factory=list)


class OccupancyMap:
    """Immutable snapshot of user-sourced cells at session start."""

    def __init__(
        self,
        user_cells: frozenset[tuple[str, str]],
        user_merges: frozenset[tuple[str, str]] = frozenset(),
    ) -> None:
        self._user_cells = user_cells
        self._user_merges = user_merges

    # --- construction ---

    @classmethod
    def empty(cls) -> OccupancyMap:
        """For sessions with no uploaded original — every cell is free."""
        return cls(frozenset())

    @classmethod
    def from_original(cls, workbook: Workbook) -> OccupancyMap:
        """Scan every sheet/cell of an uploaded workbook; record occupancy."""
        user_cells: set[tuple[str, str]] = set()
        user_merges: set[tuple[str, str]] = set()

        pyxl = workbook.raw
        for sheet_name in pyxl.sheetnames:
            ws = pyxl[sheet_name]

            # Merges: every cell inside a merge range is occupied, even if
            # openpyxl stores the spanned cells as empty shells.
            for merged in ws.merged_cells.ranges:
                user_merges.add((sheet_name, str(merged)))
                min_col, min_row, max_col, max_row = range_boundaries(str(merged))
                for r in range(min_row, max_row + 1):
                    for c in range(min_col, max_col + 1):
                        user_cells.add((sheet_name, f"{get_column_letter(c)}{r}"))

            for row in ws.iter_rows():
                for cell in row:
                    if cls._is_occupied_cell(cell):
                        user_cells.add((sheet_name, cell.coordinate))

        return cls(frozenset(user_cells), frozenset(user_merges))

    @staticmethod
    def _is_occupied_cell(cell: object) -> bool:
        """Apply the occupancy rule to a single openpyxl cell.

        - value not None covers literals and formulas (openpyxl stores
          formulas as strings starting with '=').
        - has_style covers any non-default font/fill/border/number-format.
          openpyxl sets this True on the first style assignment, which is
          occasionally a false positive but always safe: it errs toward
          "ask before overwriting", which is the correct default when
          in doubt.
        """
        value = getattr(cell, "value", None)
        if value is not None:
            return True
        if getattr(cell, "has_style", False):
            return True
        return False

    # --- queries ---

    def is_user_occupied(self, sheet: str, ref: str) -> bool:
        return (sheet, ref) in self._user_cells

    @property
    def user_cell_count(self) -> int:
        return len(self._user_cells)

    def check(self, action: Action) -> OccupancyCheck:
        """Would applying this action overwrite user-sourced data?"""
        targets = self._action_cells(action)
        overwrites = [cell for cell in targets if cell in self._user_cells]
        if not overwrites:
            return OccupancyCheck(blocked=False)
        return OccupancyCheck(
            blocked=True,
            reason=self._format_reason(overwrites),
            overwrites=overwrites,
        )

    # --- internals ---

    @staticmethod
    def _action_cells(action: Action) -> list[tuple[str, str]]:
        """Enumerate the (sheet, ref) cells an action would write.

        Only WriteFormula, WriteValue, and FormatCell target specific
        cells. AddSheet creates a new sheet; SetColumnWidth changes column
        metadata, not cell contents — both bypass the occupancy check.
        """
        if not isinstance(action, (WriteFormula, WriteValue, FormatCell)):
            return []

        sheet = action.sheet
        target = action.target
        if ":" not in target:
            return [(sheet, target)]

        min_col, min_row, max_col, max_row = range_boundaries(target)
        cells: list[tuple[str, str]] = []
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cells.append((sheet, f"{get_column_letter(c)}{r}"))
        return cells

    @staticmethod
    def _format_reason(overwrites: list[tuple[str, str]]) -> str:
        if len(overwrites) == 1:
            sheet, ref = overwrites[0]
            return f"would overwrite user-sourced cell {sheet}!{ref}"
        shown = ", ".join(f"{s}!{r}" for s, r in overwrites[:5])
        extra = f" and {len(overwrites) - 5} more" if len(overwrites) > 5 else ""
        return f"would overwrite user-sourced cells: {shown}{extra}"
