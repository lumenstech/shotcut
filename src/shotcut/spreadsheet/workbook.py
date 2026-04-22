"""Openpyxl-backed workbook that applies Action objects and delegates
formula evaluation to a lazily-constructed `Engine`.

Scope: MVP. Cell mutations flow through this class; reading workbook
state back out happens via `summary()` (for LLM context) and
`evaluate_cell()` (for verifier checks and downstream agents).

Formula *syntax* validation lives in `validator.py`. Formula *evaluation*
lives in `engine.py`.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries

from shotcut.spreadsheet.actions import (
    Action,
    AddSheet,
    FormatCell,
    SetColumnWidth,
    WriteFormula,
    WriteValue,
)

if TYPE_CHECKING:
    from shotcut.spreadsheet.engine import Engine


@dataclass
class CellSnapshot:
    value: Any
    formula: str | None
    number_format: str | None


class Workbook:
    def __init__(self, wb: OpenpyxlWorkbook):
        self._wb = wb
        self._engine: Engine | None = None

    @classmethod
    def from_xlsx(cls, path: Path) -> Workbook:
        """Load a workbook from an .xlsx file preserving formulas and structure.

        `data_only=False` keeps formulas as formulas rather than replacing
        them with their cached evaluated values — a silent-data-loss trap
        if openpyxl's default ever flips. Callers who want rich metadata
        (named ranges, merges, validations) should use `parser.parse()`
        instead; this classmethod is the minimal constructor.
        """
        return cls(load_workbook(path, data_only=False, keep_vba=True, keep_links=True))

    @classmethod
    def blank(cls) -> Workbook:
        wb = OpenpyxlWorkbook()
        # Openpyxl starts with one default sheet named "Sheet".
        return cls(wb)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self._wb.save(path)

    @property
    def raw(self) -> OpenpyxlWorkbook:
        """Underlying openpyxl workbook. Engine needs this for state mirroring."""
        return self._wb

    # --- context summary for agents ---

    def summary(self, max_cells_per_sheet: int = 50) -> dict[str, Any]:
        """Compact JSON-friendly snapshot for prompt context."""
        sheets: list[dict[str, Any]] = []
        for name in self._wb.sheetnames:
            ws = self._wb[name]
            cells: list[dict[str, Any]] = []
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells.append(
                        {
                            "ref": cell.coordinate,
                            "value": cell.value,
                            "number_format": cell.number_format or None,
                        }
                    )
                    if len(cells) >= max_cells_per_sheet:
                        break
                if len(cells) >= max_cells_per_sheet:
                    break
            sheets.append(
                {
                    "name": name,
                    "dimensions": ws.dimensions,
                    "cells": cells,
                    "truncated": len(cells) >= max_cells_per_sheet,
                }
            )
        return {"sheets": sheets}

    # --- formula evaluation ---

    def evaluate_cell(self, sheet: str, ref: str) -> Any:
        """Evaluate the cell at `sheet!ref`.

        See `shotcut.spreadsheet.engine.Engine.evaluate_cell` for error
        semantics. Evaluation is lazy: the engine materializes on first
        call and re-syncs after any `apply()`.
        """
        return self._get_engine().evaluate_cell(sheet, ref)

    def _get_engine(self) -> Engine:
        if self._engine is None:
            # Deferred import to keep import-time dependencies light for
            # code paths (e.g. audit replay) that never evaluate formulas.
            from shotcut.spreadsheet.engine import Engine

            self._engine = Engine(self)
        return self._engine

    # --- apply actions ---

    def apply(self, action: Action) -> dict[str, Any] | None:
        """Apply an action and return the previous value snapshot (for audit)."""
        result = self._dispatch(action)
        if self._engine is not None:
            self._engine.mark_dirty()
        return result

    def _dispatch(self, action: Action) -> dict[str, Any] | None:
        if isinstance(action, WriteFormula):
            return self._write(action.sheet, action.target, action.formula)
        if isinstance(action, WriteValue):
            return self._write(action.sheet, action.target, action.value)
        if isinstance(action, FormatCell):
            return self._format(action)
        if isinstance(action, AddSheet):
            self._wb.create_sheet(action.sheet)
            return None
        if isinstance(action, SetColumnWidth):
            ws = self._wb[action.sheet]
            col_letter = self._first_column_letter(action.target)
            prev = ws.column_dimensions[col_letter].width
            ws.column_dimensions[col_letter].width = action.width
            return {"width": prev}
        raise ValueError(f"Unsupported action: {action!r}")

    # --- internals ---

    def _write(self, sheet: str, target: str, value: Any) -> dict[str, Any]:
        ws = self._wb[sheet]
        previous: dict[str, Any] = {}
        if ":" in target:
            min_col, min_row, max_col, max_row = range_boundaries(target)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    cell = ws.cell(row=row, column=col)
                    previous[cell.coordinate] = cell.value
                    cell.value = value
        else:
            col, row = coordinate_from_string(target)
            cell = ws[f"{col}{row}"]
            previous[cell.coordinate] = cell.value
            cell.value = value
        return {"cells": previous}

    def _format(self, action: FormatCell) -> dict[str, Any]:
        ws = self._wb[action.sheet]
        previous: dict[str, Any] = {}

        if ":" in action.target:
            min_col, min_row, max_col, max_row = range_boundaries(action.target)
            cells = [
                ws.cell(row=r, column=c)
                for r in range(min_row, max_row + 1)
                for c in range(min_col, max_col + 1)
            ]
        else:
            cells = [ws[action.target]]

        for cell in cells:
            previous[cell.coordinate] = {
                "number_format": cell.number_format,
                "font": {"bold": cell.font.bold, "italic": cell.font.italic},
            }
            if action.number_format is not None:
                cell.number_format = action.number_format
            if action.bold is not None or action.italic is not None:
                cell.font = Font(
                    bold=action.bold if action.bold is not None else cell.font.bold,
                    italic=action.italic if action.italic is not None else cell.font.italic,
                )
        return previous

    @staticmethod
    def _first_column_letter(target: str) -> str:
        # openpyxl has no type stubs, so these helpers return Any; cast at
        # the boundary to keep our internal types sharp.
        if ":" in target:
            min_col, _, _, _ = range_boundaries(target)
            return str(get_column_letter(min_col))
        col, _ = coordinate_from_string(target)
        return str(col)
