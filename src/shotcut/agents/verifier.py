"""Five-level workbook verifier.

Levels, run in order:

1. **Syntax**    — every formula parses (existing `validator.validate_formula`).
2. **Reference** — all refs resolve: sheets exist, cross-sheet refs valid,
   named ranges defined. Surfaces as `#REF!` or `#NAME?` from the engine.
3. **Cycle**     — dependency graph acyclic. Surfaces as `#CIRC!`.
4. **Numerical** — evaluate via Stage 1 engine; flag `#DIV/0!`, `#VALUE!`,
   `#N/A`, `#NUM!`, NaN, infinity.
5. **Semantic**  — Opus 4.7 with the cached prompt in `verifier_prompts.py`.
   Runs only if levels 1–4 produced no critical findings — there's no
   point asking a model to reason about a workbook we already know is
   broken.

Pending-approval handling (Stage 3 amendment). The verifier evaluates a
*clone* of the workbook with every `pending_approval` action from the
current turn applied. The persisted workbook is never mutated. Findings
whose target cell falls inside a pending action's range are attributed
back to that action so the orchestrator can append the finding to the
action's `reasoning` column. The action's status stays
`pending_approval` — verifier never auto-rejects.

The clone-and-discard approach is a deliberately cheap Stage 4 shape.
The copy-on-write `Workbook.with_actions_applied` upgrade is noted in
`docs/decisions/0002-schema-consolidation.md` as the Stage 6 follow-up.
"""
from __future__ import annotations

import enum
import io
import json
import math
import uuid
from collections import defaultdict
from collections.abc import Iterator
from typing import TYPE_CHECKING

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from shotcut.agents.verifier_prompts import SEMANTIC_SYSTEM_PROMPT
from shotcut.config import settings
from shotcut.llm.client import cached_system, get_client
from shotcut.spreadsheet.engine import (
    CircularReferenceError,
    DivisionByZeroError,
    FormulaEvaluationError,
    InvalidReferenceError,
    InvalidValueError,
    NotAvailableError,
    NumericError,
    UnsupportedFormulaError,
)
from shotcut.spreadsheet.validator import validate_formula
from shotcut.spreadsheet.workbook import Workbook, _release_empty_vba_archive

if TYPE_CHECKING:
    from shotcut.spreadsheet.actions import Action as AgentAction


# ---------------------------------------------------------------------------
# Public types
# ---------------------------------------------------------------------------


class VerifierLevel(str, enum.Enum):
    SYNTAX = "syntax"
    REFERENCE = "reference"
    CYCLE = "cycle"
    NUMERICAL = "numerical"
    SEMANTIC = "semantic"


class VerifierSeverity(str, enum.Enum):
    CRITICAL = "critical"
    WARNING = "warning"
    INFO = "info"


class VerifierFinding(BaseModel):
    level: VerifierLevel
    severity: VerifierSeverity
    sheet: str | None = None
    cell: str | None = None
    message: str
    suggested_fix: str | None = None


class VerificationReport(BaseModel):
    findings: list[VerifierFinding]
    confidence: float = Field(ge=0.0, le=1.0)

    @property
    def critical(self) -> list[VerifierFinding]:
        return [f for f in self.findings if f.severity is VerifierSeverity.CRITICAL]

    @property
    def has_critical(self) -> bool:
        return any(f.severity is VerifierSeverity.CRITICAL for f in self.findings)


# Intermediate shape used by `_semantic_pass` for messages.parse(). Kept
# separate from `VerificationReport` so the LLM returns a narrower schema.
class _SemanticFindings(BaseModel):
    findings: list[VerifierFinding]
    confidence: float = Field(ge=0.0, le=1.0, default=1.0)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


async def verify(
    workbook: Workbook,
    *,
    prompt: str,
    pending_actions: list[AgentAction] | None = None,
) -> VerificationReport:
    """Run all five levels against a clone of the workbook with pending
    actions applied.

    The caller's `workbook` is not mutated. `pending_actions` is the list
    of `Action` domain objects for pending_approval audit rows in the
    current turn; they're applied to the clone before verification so the
    user can see the findings the applied state *would* produce.
    """
    target = _clone_with_pending(workbook, pending_actions or [])

    findings: list[VerifierFinding] = []
    findings.extend(_syntax_pass(target))
    findings.extend(_engine_pass(target))

    # Short-circuit the semantic (LLM-billed) level when earlier levels
    # already found critical issues. No point asking Opus to reason about
    # a workbook with #REF! everywhere.
    has_critical = any(f.severity is VerifierSeverity.CRITICAL for f in findings)
    if not has_critical:
        semantic = await _semantic_pass(target, prompt)
        findings.extend(semantic.findings)
        confidence = semantic.confidence
    else:
        # Deterministic levels saw a critical issue; we're confident in
        # reporting it but skipped the LLM pass, so lower nominal score.
        confidence = 0.5

    return VerificationReport(findings=findings, confidence=confidence)


def attribute_to_actions(
    findings: list[VerifierFinding],
    pending_actions: list[AgentAction],
    row_id_by_client_id: dict[uuid.UUID, uuid.UUID],
) -> dict[uuid.UUID, list[VerifierFinding]]:
    """Map each finding to the pending Action (by audit row id) whose
    target range contains the finding's cell.

    Keyed on the domain Action's `client_action_id`, which is stable
    across serialization and Stage 5 branches. Python `id()` is NOT
    stable across those boundaries and was replaced by the Stage 5
    erratum (see docs/decisions/0002-schema-consolidation.md).

    `row_id_by_client_id` is built by the caller from the paired list
    of (domain_action, db_row): `{a.client_action_id: row.id for ...}`.
    """
    out: dict[uuid.UUID, list[VerifierFinding]] = defaultdict(list)
    for finding in findings:
        if finding.sheet is None or finding.cell is None:
            continue
        for action in pending_actions:
            if _action_touches(action, finding.sheet, finding.cell):
                row_id = row_id_by_client_id.get(action.client_action_id)
                if row_id is not None:
                    out[row_id].append(finding)
                break
    return dict(out)


# ---------------------------------------------------------------------------
# Level 1: Syntax
# ---------------------------------------------------------------------------


def _syntax_pass(workbook: Workbook) -> list[VerifierFinding]:
    findings: list[VerifierFinding] = []
    for sheet_name, ref, value in _iter_formula_cells(workbook):
        for issue in validate_formula(value):
            severity = (
                VerifierSeverity.CRITICAL if issue.severity == "error"
                else VerifierSeverity.WARNING
            )
            findings.append(
                VerifierFinding(
                    level=VerifierLevel.SYNTAX,
                    severity=severity,
                    sheet=sheet_name,
                    cell=ref,
                    message=f"{issue.message} in formula {issue.formula!r}",
                )
            )
    return findings


# ---------------------------------------------------------------------------
# Levels 2-4: Reference / Cycle / Numerical (engine-driven)
# ---------------------------------------------------------------------------


def _engine_pass(workbook: Workbook) -> list[VerifierFinding]:
    """Evaluate every formula cell and classify the resulting error kind."""
    findings: list[VerifierFinding] = []
    for sheet_name, ref, _formula in _iter_formula_cells(workbook):
        try:
            value = workbook.evaluate_cell(sheet_name, ref)
        except CircularReferenceError as exc:
            findings.append(_engine_finding(VerifierLevel.CYCLE, VerifierSeverity.CRITICAL,
                                            sheet_name, ref, str(exc)))
            continue
        except InvalidReferenceError as exc:
            findings.append(_engine_finding(VerifierLevel.REFERENCE, VerifierSeverity.CRITICAL,
                                            sheet_name, ref, str(exc)))
            continue
        except UnsupportedFormulaError as exc:
            # Unsupported functions manifest as #NAME? — reference-ish.
            findings.append(_engine_finding(VerifierLevel.REFERENCE, VerifierSeverity.WARNING,
                                            sheet_name, ref, str(exc),
                                            suggested_fix="replace with a supported function"))
            continue
        except DivisionByZeroError as exc:
            findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.WARNING,
                                            sheet_name, ref, str(exc),
                                            suggested_fix="wrap in IFERROR or guard the denominator"))
            continue
        except NotAvailableError as exc:
            findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.WARNING,
                                            sheet_name, ref, str(exc)))
            continue
        except (InvalidValueError, NumericError) as exc:
            findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.CRITICAL,
                                            sheet_name, ref, str(exc)))
            continue
        except FormulaEvaluationError as exc:
            # Fallback for any unclassified engine error.
            findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.WARNING,
                                            sheet_name, ref, str(exc)))
            continue

        # NaN / infinity don't raise — they come through as float values.
        if isinstance(value, float):
            if math.isnan(value):
                findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.WARNING,
                                                sheet_name, ref, "result is NaN"))
            elif math.isinf(value):
                findings.append(_engine_finding(VerifierLevel.NUMERICAL, VerifierSeverity.WARNING,
                                                sheet_name, ref, "result is infinity"))
    return findings


def _engine_finding(
    level: VerifierLevel,
    severity: VerifierSeverity,
    sheet: str,
    ref: str,
    message: str,
    suggested_fix: str | None = None,
) -> VerifierFinding:
    return VerifierFinding(
        level=level, severity=severity, sheet=sheet, cell=ref,
        message=message, suggested_fix=suggested_fix,
    )


# ---------------------------------------------------------------------------
# Level 5: Semantic (LLM)
# ---------------------------------------------------------------------------


async def _semantic_pass(workbook: Workbook, prompt: str) -> _SemanticFindings:
    client = get_client()
    summary = workbook.summary(max_cells_per_sheet=500)
    user_message = (
        f"Original request:\n{prompt}\n\n"
        f"Workbook state (JSON):\n{json.dumps(summary, default=str)}"
    )
    response = await client.messages.parse(
        model=settings.verifier_model,
        max_tokens=16000,
        thinking={"type": "adaptive"},
        output_config={"effort": "high"},
        system=cached_system(SEMANTIC_SYSTEM_PROMPT),
        messages=[{"role": "user", "content": user_message}],
        output_format=_SemanticFindings,
    )
    result = response.parsed_output
    if result is None:
        raise RuntimeError("verifier: semantic level returned no parseable output")
    # Force the level on findings — the LLM can hallucinate; we label
    # ourselves.
    for f in result.findings:
        f.level = VerifierLevel.SEMANTIC
    return result


# ---------------------------------------------------------------------------
# Internals
# ---------------------------------------------------------------------------


def _iter_formula_cells(workbook: Workbook) -> Iterator[tuple[str, str, str]]:
    """Yield (sheet, ref, formula_str) for every cell whose value is a
    formula string (starts with '=')."""
    pyxl = workbook.raw
    for sheet_name in pyxl.sheetnames:
        ws = pyxl[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    yield sheet_name, cell.coordinate, value


def _clone_with_pending(
    workbook: Workbook, pending: list[AgentAction]
) -> Workbook:
    """Serialize the workbook to BytesIO and reload, then apply each
    pending action to the clone. Original is not mutated.

    Cost: one full openpyxl round-trip per verifier call. Acceptable for
    MVP workbook sizes; Stage 6's checkpoint replay motivates upgrading
    to `Workbook.with_actions_applied` (see 0002 decision doc).
    """
    buffer = io.BytesIO()
    workbook.raw.save(buffer)
    buffer.seek(0)
    pyxl = load_workbook(buffer, data_only=False, keep_vba=True, keep_links=True)
    _release_empty_vba_archive(pyxl)
    clone = Workbook(pyxl)
    for action in pending:
        clone.apply(action)
    return clone


def _action_touches(action: AgentAction, sheet: str, ref: str) -> bool:
    """Whether `action` writes to `sheet!ref`. Reuses OccupancyMap's
    cell-enumeration helper to avoid duplicating the range-expansion logic."""
    from shotcut.spreadsheet.occupancy import OccupancyMap

    return (sheet, ref) in OccupancyMap._action_cells(action)
